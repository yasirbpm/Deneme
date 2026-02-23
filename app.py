import argparse
import re
import time
from pathlib import Path
from typing import Dict, List, Optional, Set
from urllib.parse import quote

import pandas as pd
from openpyxl.utils import get_column_letter
from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
from playwright.sync_api import sync_playwright


BASE_MAPS_SEARCH_URL = "https://www.google.com/maps/search/"


def load_keywords_from_txt(path: Optional[str]) -> List[str]:
    if not path:
        return []

    file_path = Path(path)
    if not file_path.exists():
        raise FileNotFoundError(f"Keyword dosyası bulunamadı: {path}")

    keywords: List[str] = []
    for line in file_path.read_text(encoding="utf-8").splitlines():
        normalized = line.strip()
        if normalized and not normalized.startswith("#"):
            keywords.append(normalized)

    return keywords


def build_queries(
    city: str,
    country: str,
    category: Optional[str],
    direct_query: Optional[str],
    keywords: List[str],
) -> List[str]:
    base_queries: List[str] = []

    if direct_query:
        base_queries.append(direct_query.strip())

    if category:
        base_queries.append(category.strip())

    base_queries.extend(keywords)

    if not base_queries:
        raise ValueError("En az bir sorgu kaynağı verin: --query veya --category veya --keywords-file")

    final_queries: List[str] = []
    seen: Set[str] = set()

    for item in base_queries:
        normalized = " ".join(item.split()).strip()
        if not normalized:
            continue
        full_query = f"{normalized} in {city}, {country}"
        key = full_query.lower()
        if key not in seen:
            seen.add(key)
            final_queries.append(full_query)

    return final_queries


def parse_rating(aria_label: str) -> Optional[float]:
    if not aria_label:
        return None
    m = re.search(r"([0-9]+(?:[\.,][0-9]+)?)", aria_label)
    if not m:
        return None
    return float(m.group(1).replace(",", "."))


def extract_first_text(page, selectors: List[str]) -> Optional[str]:
    for selector in selectors:
        locator = page.locator(selector).first
        if locator.count() > 0:
            text = locator.inner_text().strip()
            if text:
                return text
    return None


def extract_place_id_from_url(url: str) -> Optional[str]:
    if not url:
        return None

    m = re.search(r"!1s([^!]+)", url)
    if m:
        return m.group(1)

    m2 = re.search(r"/place/([^/]+)/", url)
    if m2:
        return m2.group(1)

    return None


def should_exclude(name: str, category_text: str, exclude_terms: List[str]) -> bool:
    if not exclude_terms:
        return False

    haystack = f"{name} {category_text}".lower()
    for term in exclude_terms:
        t = term.strip().lower()
        if t and t in haystack:
            return True
    return False


def scrape_business_details(page, source_query: str, exclude_terms: List[str]) -> Optional[Dict]:
    try:
        page.wait_for_selector("h1", timeout=10000)
    except PlaywrightTimeoutError:
        return None

    name = extract_first_text(page, ["h1.DUwDvf", "h1"])
    if not name:
        return None

    category_text = extract_first_text(
        page,
        [
            "button[jsaction*='pane.rating.category']",
            "button.DkEaL",
            "div[role='main'] button:has-text('·')",
        ],
    ) or ""

    if should_exclude(name, category_text, exclude_terms):
        return None

    address = extract_first_text(page, ["button[data-item-id='address']", "div[data-item-id='address']"])
    phone = extract_first_text(page, ["button[data-item-id^='phone']", "div[data-item-id^='phone']"])

    website = None
    website_loc = page.locator("a[data-item-id='authority']").first
    if website_loc.count() > 0:
        href = website_loc.get_attribute("href")
        if href:
            website = href.strip()

    rating = None
    rating_loc = page.locator("span[role='img'][aria-label*='yıldız'], span[role='img'][aria-label*='star']").first
    if rating_loc.count() > 0:
        aria_label = rating_loc.get_attribute("aria-label") or ""
        rating = parse_rating(aria_label)

    reviews_text = extract_first_text(
        page,
        [
            "button[jsaction*='pane.reviewChart.moreReviews']",
            "button[aria-label*='yorum']",
            "button[aria-label*='review']",
        ],
    )
    user_ratings_total = None
    if reviews_text:
        m = re.search(r"([0-9\.,]+)", reviews_text)
        if m:
            user_ratings_total = int(m.group(1).replace(".", "").replace(",", ""))

    current_url = page.url
    place_id = extract_place_id_from_url(current_url)

    return {
        "source_query": source_query,
        "place_id": place_id,
        "name": name,
        "address": address,
        "phone": phone,
        "website": website,
        "has_website": bool(website and website.strip()),
        "rating": rating,
        "user_ratings_total": user_ratings_total,
        "types": category_text,
        "business_status": None,
        "google_maps_url": current_url,
    }


def scrape_query(
    page,
    query: str,
    max_results_per_query: int,
    exclude_terms: List[str],
    seen_place_ids: Set[str],
) -> List[Dict]:
    url = f"{BASE_MAPS_SEARCH_URL}{quote(query)}"
    page.goto(url, wait_until="domcontentloaded", timeout=60000)

    results: List[Dict] = []
    feed = page.locator("div[role='feed']")
    if feed.count() == 0:
        return results

    scrollable = feed.first
    no_change_rounds = 0
    last_count = 0

    while len(results) < max_results_per_query and no_change_rounds < 8:
        cards = page.locator("a[href*='/maps/place/']")
        count = cards.count()

        for i in range(count):
            if len(results) >= max_results_per_query:
                break

            card = cards.nth(i)
            href = card.get_attribute("href")
            if not href:
                continue

            place_id_hint = extract_place_id_from_url(href) or href
            if place_id_hint in seen_place_ids:
                continue

            try:
                card.click(timeout=5000)
                time.sleep(1.5)
            except PlaywrightTimeoutError:
                continue

            details = scrape_business_details(page, query, exclude_terms)
            if not details:
                continue

            real_place_id = details.get("place_id") or place_id_hint
            if real_place_id in seen_place_ids:
                continue

            details["place_id"] = real_place_id
            seen_place_ids.add(real_place_id)
            results.append(details)

        scrollable.evaluate("el => el.scrollBy(0, el.scrollHeight)")
        time.sleep(1.5)

        new_count = page.locator("a[href*='/maps/place/']").count()
        if new_count == last_count:
            no_change_rounds += 1
        else:
            no_change_rounds = 0
            last_count = new_count

    return results


def collect_businesses(
    queries: List[str],
    max_results_per_query: int,
    exclude_terms: List[str],
    headless: bool,
) -> pd.DataFrame:
    rows: List[Dict] = []
    seen_place_ids: Set[str] = set()

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=headless)
        context = browser.new_context(locale="tr-TR")
        page = context.new_page()

        for query in queries:
            query_rows = scrape_query(
                page=page,
                query=query,
                max_results_per_query=max_results_per_query,
                exclude_terms=exclude_terms,
                seen_place_ids=seen_place_ids,
            )
            rows.extend(query_rows)

        context.close()
        browser.close()

    df = pd.DataFrame(rows)
    if not df.empty:
        df = df.sort_values(by=["has_website", "rating", "user_ratings_total"], ascending=[True, False, False])
    return df


def apply_excel_filters_and_widths(writer: pd.ExcelWriter, sheet_name: str, df: pd.DataFrame) -> None:
    ws = writer.sheets[sheet_name]
    if df.empty:
        return

    max_col = df.shape[1]
    max_row = df.shape[0] + 1
    ws.auto_filter.ref = f"A1:{get_column_letter(max_col)}{max_row}"

    for idx, column in enumerate(df.columns, start=1):
        max_len = max(df[column].astype(str).map(len).max(), len(column))
        ws.column_dimensions[get_column_letter(idx)].width = min(max_len + 2, 60)



def export_excel(df: pd.DataFrame, output_path: str) -> None:
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)

    no_website_df = df[df["has_website"] == False].copy() if not df.empty else df.copy()  # noqa: E712

    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="all_businesses", index=False)
        no_website_df.to_excel(writer, sheet_name="no_website", index=False)

        apply_excel_filters_and_widths(writer, "all_businesses", df)
        apply_excel_filters_and_widths(writer, "no_website", no_website_df)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Google Maps scraping ile işletme verisi toplayıp Excel'e aktarır.")

    parser.add_argument("--country", required=True, help="Ülke adı (örn: Turkey)")
    parser.add_argument("--city", required=True, help="Şehir adı (örn: Istanbul)")

    parser.add_argument("--category", help="Kategori (örn: hotel, dental clinic, restaurant)")
    parser.add_argument("--query", help="Serbest sorgu (kategori yerine veya kategoriyle birlikte kullanılabilir)")
    parser.add_argument("--keywords-file", help="Satır satır keyword içeren txt dosyası")

    parser.add_argument(
        "--exclude-categories",
        default="",
        help="Virgülle ayrılmış hariç tutulacak kategori/terimler (örn: bar,night_club,casino)",
    )

    parser.add_argument("--max-results-per-query", type=int, default=60, help="Her sorgu için maksimum işletme")
    parser.add_argument("--output", default="output/businesses.xlsx", help="Excel çıktı dosyası")
    parser.add_argument("--headed", action="store_true", help="Tarayıcıyı görünür modda aç")
    return parser.parse_args()


def main() -> None:
    args = parse_args()

    keywords = load_keywords_from_txt(args.keywords_file)
    exclude_terms = [item.strip() for item in args.exclude_categories.split(",") if item.strip()]

    queries = build_queries(
        city=args.city,
        country=args.country,
        category=args.category,
        direct_query=args.query,
        keywords=keywords,
    )

    df = collect_businesses(
        queries=queries,
        max_results_per_query=args.max_results_per_query,
        exclude_terms=exclude_terms,
        headless=not args.headed,
    )

    export_excel(df, args.output)
    total = len(df)
    no_website_count = int((~df["has_website"]).sum()) if total else 0

    print(f"Konum: {args.city}, {args.country}")
    print(f"Sorgu sayısı: {len(queries)}")
    print(f"Toplam işletme: {total}")
    print(f"Websitesi olmayan işletme: {no_website_count}")
    print(f"Excel yazıldı: {args.output}")


if __name__ == "__main__":
    main()
